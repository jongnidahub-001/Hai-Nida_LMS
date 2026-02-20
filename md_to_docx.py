#!/usr/bin/env python3
"""
สคริปต์แปลงไฟล์ Markdown (.md) เป็นไฟล์ Word (.docx)
รองรับการแปลงรูปแบบต่างๆ รวมถึงรูปภาพ
"""

import os
import sys
import argparse
from pathlib import Path

try:
    import pypandoc
    PYPANDOC_AVAILABLE = True
except ImportError:
    PYPANDOC_AVAILABLE = False
    print("Warning: pypandoc not available. Trying alternative method...")

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import markdown
    from markdown.extensions import codehilite, tables, fenced_code
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def convert_with_pypandoc(input_file, output_file):
    """แปลงไฟล์ markdown เป็น docx โดยใช้ pypandoc"""
    try:
        pypandoc.convert_file(
            input_file,
            'docx',
            outputfile=output_file,
            extra_args=['--standalone', '--toc']
        )
        print(f"✓ แปลงไฟล์สำเร็จ: {output_file}")
        return True
    except Exception as e:
        print(f"✗ เกิดข้อผิดพลาดในการแปลงด้วย pypandoc: {str(e)}")
        return False


def markdown_to_html(md_content):
    """แปลง markdown เป็น HTML"""
    extensions = [
        'markdown.extensions.codehilite',
        'markdown.extensions.tables',
        'markdown.extensions.fenced_code',
        'markdown.extensions.toc',
        'markdown.extensions.extra'
    ]
    html = markdown.markdown(md_content, extensions=extensions)
    return html


def html_to_docx(html_content, output_file, input_file_path):
    """แปลง HTML เป็น DOCX"""
    from bs4 import BeautifulSoup
    import re
    
    doc = Document()
    soup = BeautifulSoup(html_content, 'html.parser')
    
    # ตั้งค่า style เริ่มต้น
    style = doc.styles['Normal']
    font = style.font
    font.name = 'TH Sarabun New'
    font.size = Pt(14)
    
    # ฟังก์ชันแปลง HTML elements เป็น docx
    def process_element(element):
        if element.name is None:  # Text node
            text = str(element).strip()
            if text:
                return text
            return None
        
        tag = element.name.lower()
        
        # หัวข้อ
        if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            level = int(tag[1])
            heading = doc.add_heading(level=level)
            for child in element.children:
                text = process_element(child)
                if text:
                    heading.add_run(text)
            return None
        
        # ย่อหน้า
        elif tag == 'p':
            p = doc.add_paragraph()
            for child in element.children:
                text = process_element(child)
                if text:
                    p.add_run(text)
            return None
        
        # รายการ
        elif tag in ['ul', 'ol']:
            for li in element.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Bullet' if tag == 'ul' else 'List Number')
                for child in li.children:
                    text = process_element(child)
                    if text:
                        p.add_run(text)
            return None
        
        # รูปภาพ
        elif tag == 'img':
            src = element.get('src', '')
            if src:
                # ลองหาไฟล์รูปภาพ
                img_path = None
                if os.path.isabs(src):
                    img_path = src
                else:
                    # ลองหาในโฟลเดอร์เดียวกับไฟล์ input
                    base_dir = os.path.dirname(input_file_path)
                    img_path = os.path.join(base_dir, src)
                    if not os.path.exists(img_path):
                        # ลองหาในโฟลเดอร์ images
                        img_path = os.path.join(base_dir, 'images', os.path.basename(src))
                
                if img_path and os.path.exists(img_path):
                    try:
                        doc.add_picture(img_path, width=Inches(5))
                    except Exception as e:
                        print(f"Warning: ไม่สามารถเพิ่มรูปภาพ {src}: {str(e)}")
                else:
                    print(f"Warning: ไม่พบไฟล์รูปภาพ: {src}")
            return None
        
        # ตาราง
        elif tag == 'table':
            table = doc.add_table(rows=0, cols=0)
            table.style = 'Light Grid Accent 1'
            
            # หา header
            thead = element.find('thead')
            if thead:
                header_row = thead.find('tr')
                if header_row:
                    cells = header_row.find_all(['th', 'td'])
                    if cells:
                        table.add_row()
                        for i, cell in enumerate(cells):
                            if i == 0:
                                table.columns[i].width = Inches(1.5)
                            table.rows[0].cells[i].text = cell.get_text().strip()
            
            # หา body
            tbody = element.find('tbody') or element
            for tr in tbody.find_all('tr', recursive=False):
                if tr.find('th'):  # Skip header row if already processed
                    continue
                cells = tr.find_all('td')
                if cells:
                    row = table.add_row()
                    for i, cell in enumerate(cells):
                        if i < len(row.cells):
                            row.cells[i].text = cell.get_text().strip()
            return None
        
        # Code blocks
        elif tag in ['pre', 'code']:
            if tag == 'pre':
                code = element.find('code')
                if code:
                    p = doc.add_paragraph(style='No Spacing')
                    run = p.add_run(code.get_text())
                    run.font.name = 'Consolas'
                    run.font.size = Pt(10)
            return None
        
        # Bold, Italic
        elif tag == 'strong' or tag == 'b':
            return element.get_text()
        elif tag == 'em' or tag == 'i':
            return element.get_text()
        
        # Link
        elif tag == 'a':
            return element.get_text()
        
        # อื่นๆ - ดึงข้อความ
        else:
            return element.get_text()
    
    # ประมวลผลเนื้อหา
    for element in soup.children:
        process_element(element)
    
    # บันทึกไฟล์
    doc.save(output_file)
    print(f"✓ แปลงไฟล์สำเร็จ: {output_file}")


def convert_with_alternative(input_file, output_file):
    """แปลงไฟล์ markdown เป็น docx โดยใช้วิธีอื่น (markdown -> html -> docx)"""
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        print("✗ ต้องการ BeautifulSoup4: pip install beautifulsoup4")
        return False
    
    # อ่านไฟล์ markdown
    with open(input_file, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # แปลง markdown เป็น HTML
    html_content = markdown_to_html(md_content)
    
    # แปลง HTML เป็น DOCX
    html_to_docx(html_content, output_file, input_file)
    
    return True


def main():
    parser = argparse.ArgumentParser(
        description='แปลงไฟล์ Markdown (.md) เป็นไฟล์ Word (.docx)',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
ตัวอย่างการใช้งาน:
  python md_to_docx.py input.md
  python md_to_docx.py input.md -o output.docx
  python md_to_docx.py input.md --output output.docx
        """
    )
    
    parser.add_argument(
        'input_file',
        type=str,
        help='ไฟล์ Markdown ที่ต้องการแปลง (.md)'
    )
    
    parser.add_argument(
        '-o', '--output',
        type=str,
        default=None,
        help='ไฟล์ output (.docx) ถ้าไม่ระบุจะใช้ชื่อเดียวกับไฟล์ input'
    )
    
    args = parser.parse_args()
    
    # ตรวจสอบไฟล์ input
    input_path = Path(args.input_file)
    if not input_path.exists():
        print(f"✗ ไม่พบไฟล์: {args.input_file}")
        sys.exit(1)
    
    if not input_path.suffix.lower() == '.md':
        print(f"✗ ไฟล์ต้องเป็น .md: {args.input_file}")
        sys.exit(1)
    
    # กำหนดไฟล์ output
    if args.output:
        output_path = Path(args.output)
    else:
        output_path = input_path.with_suffix('.docx')
    
    # แปลงไฟล์
    print(f"กำลังแปลง: {input_path} -> {output_path}")
    
    # ลองใช้ pypandoc ก่อน (ถ้ามี)
    if PYPANDOC_AVAILABLE:
        print("ใช้ pypandoc ในการแปลง...")
        if convert_with_pypandoc(str(input_path), str(output_path)):
            sys.exit(0)
        else:
            print("ลองใช้วิธีอื่น...")
    
    # ใช้วิธีอื่น
    if DOCX_AVAILABLE:
        print("ใช้ markdown + python-docx ในการแปลง...")
        if convert_with_alternative(str(input_path), str(output_path)):
            sys.exit(0)
    else:
        print("✗ ต้องการไลบรารีเพิ่มเติม:")
        print("  pip install python-docx markdown beautifulsoup4")
        print("  หรือ")
        print("  pip install pypandoc")
        print("  (และติดตั้ง pandoc: https://pandoc.org/installing.html)")
        sys.exit(1)


if __name__ == '__main__':
    main()

